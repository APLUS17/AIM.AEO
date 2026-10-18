"""Builds the free five-question test spreadsheet and verifies it and the snippet pack.

Run from anywhere:  python3 plan/assets/make_and_verify.py [path/to/schemaorg-current-https.jsonld]

Checks:
  1. The spreadsheet's formulas recalculate correctly (LibreOffice recalculates a copy filled with sample answers).
  2. The robots.txt files behave as the post says (Python's robots.txt parser).
  3. Each JSON-LD snippet parses, uses real schema.org types, and uses only properties valid for those types.
Writes VERIFIED.md next to this script. Nothing here needs network access except the optional schema.org file.
"""
import json, os, shutil, subprocess, sys, tempfile, urllib.robotparser
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "five-question-ai-visibility-test.xlsx")
LOG_ROWS = 90  # three full rounds of 5 questions x 3 engines x 2 runs
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


def build_xlsx(path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("Five-question AI visibility test", True),
        ("", False),
        ("What this is: a way to see what ChatGPT, Claude and Perplexity say about your business. About 30 minutes, no tools.", False),
        ("", False),
        ("Steps", True),
        ("1. On the Questions sheet, fill in the yellow cells. The five questions write themselves.", False),
        ("2. Open each engine in a fresh chat. Turn off memory or personalization if the tool has it.", False),
        ("3. Paste each question with the exact wording. Use the same wording on every engine.", False),
        ("4. Ask each question at least twice. Answers vary from run to run.", False),
        ("5. Log every answer on the Log sheet: the date, the engine, the question number, and what you saw.", False),
        ("6. Read the Summary sheet. It fills in by itself.", False),
        ("", False),
        ("How to read the Summary", True),
        ("Low mention rate on the unbranded questions (1 and 2): a visibility problem. The AI does not connect you with your category.", False),
        ("Mentioned, but low in the list or described generically: a position problem. Check the Position and 'What it said' columns.", False),
        ("Wrong or outdated facts on the branded questions (3 to 5): an accuracy problem. Note exactly what was wrong and where it may have come from.", False),
        ("", False),
        ("Keep it repeatable: save the exact wording and run the same five questions once a month. One run proves little; look for a pattern.", False),
        ("The chat apps your buyers use can answer differently from the developer tools (APIs). Use the apps for this test.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110

    q = wb.create_sheet("Questions")
    q["A1"], q["A1"].font = "Fill in the yellow cells", BOLD
    inputs = [("Your brand", "B2"), ("Your category (e.g. invoicing tool)", "B3"), ("Type of buyer, with the article (e.g. a five-person agency)", "B4"),
              ("Specific need or budget", "B5"), ("Main competitor", "B6")]
    for i, (label, _) in enumerate(inputs, 2):
        q.cell(row=i, column=1, value=label)
        q.cell(row=i, column=2).fill = YELLOW
    for col, head in zip("ABC", ["#", "Type", "Question to paste"]):
        q[f"{col}8"], q[f"{col}8"].font = head, BOLD
    ask = '"Fill in the yellow cells above"'
    questions = [
        (1, "Unbranded", f'=IF(OR($B$3="",$B$4=""),{ask},"What is the best "&$B$3&" for "&$B$4&"?")'),
        (2, "Unbranded", f'=IF(OR($B$3="",$B$5=""),{ask},"What are the top "&$B$3&" options for "&$B$5&"?")'),
        (3, "Branded", f'=IF($B$2="",{ask},"What is "&$B$2&" and who is it for?")'),
        (4, "Branded", f'=IF(OR($B$2="",$B$6=""),{ask},"Is "&$B$2&" worth it compared with "&$B$6&"?")'),
        (5, "Branded", f'=IF($B$2="",{ask},"How much does "&$B$2&" cost and what is included?")'),
    ]
    for r, (n, kind, formula) in enumerate(questions, 9):
        q.cell(row=r, column=1, value=n)
        q.cell(row=r, column=2, value=kind)
        q.cell(row=r, column=3, value=formula)
    q.column_dimensions["A"].width = 38
    q.column_dimensions["B"].width = 38
    q.column_dimensions["C"].width = 90

    log = wb.create_sheet("Log")
    heads = ["Date", "Engine", "Question #", "Question type", "Run #", "Mentioned?", "Position", "What it said about you",
             "Wrong or outdated?", "Competitors named", "Sources cited"]
    for i, h in enumerate(heads, 1):
        c = log.cell(row=1, column=i, value=h)
        c.font = BOLD
    last = LOG_ROWS + 1
    for r in range(2, last + 1):
        log.cell(row=r, column=4, value=f'=IF(C{r}="","",IF(C{r}<=2,"Unbranded","Branded"))')
    for rng, items in [
        (f"B2:B{last}", '"ChatGPT,Claude,Perplexity,Other"'),
        (f"C2:C{last}", '"1,2,3,4,5"'),
        (f"F2:F{last}", '"Yes,No"'),
        (f"G2:G{last}", '"1,2,3,4,5 or lower,Not listed"'),
        (f"I2:I{last}", '"Yes,No"'),
    ]:
        dv = DataValidation(type="list", formula1=items, allow_blank=True)
        log.add_data_validation(dv)
        dv.add(rng)
    for col, w in zip("ABCDEFGHIJK", [12, 12, 11, 14, 8, 12, 14, 50, 18, 30, 30]):
        log.column_dimensions[col].width = w
    log.freeze_panes = "A2"

    s = wb.create_sheet("Summary")
    s["A1"], s["A1"].font = "Summary (fills in from the Log sheet)", Font(bold=True, size=14)
    heads = ["Engine", "Answers logged", "Unbranded answers", "Mentioned in unbranded", "Unbranded mention rate",
             "Branded answers", "Wrong or outdated facts"]
    for i, h in enumerate(heads, 1):
        c = s.cell(row=3, column=i, value=h)
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True)
    rng = lambda col: f"Log!${col}$2:${col}${last}"
    for r, eng in enumerate(["ChatGPT", "Claude", "Perplexity"], 4):
        s.cell(row=r, column=1, value=eng)
        s.cell(row=r, column=2, value=f"=COUNTIFS({rng('B')},$A{r})")
        s.cell(row=r, column=3, value=f'=COUNTIFS({rng("B")},$A{r},{rng("D")},"Unbranded")')
        s.cell(row=r, column=4, value=f'=COUNTIFS({rng("B")},$A{r},{rng("D")},"Unbranded",{rng("F")},"Yes")')
        s.cell(row=r, column=5, value=f'=IF(C{r}=0,"",D{r}/C{r})').number_format = "0%"
        s.cell(row=r, column=6, value=f'=COUNTIFS({rng("B")},$A{r},{rng("D")},"Branded")')
        s.cell(row=r, column=7, value=f'=COUNTIFS({rng("B")},$A{r},{rng("I")},"Yes")')
    s["A7"], s["A7"].font = "All engines", BOLD
    for col in "BCDFG":
        s[f"{col}7"] = f"=SUM({col}4:{col}6)"
    s["E7"] = '=IF(C7=0,"",D7/C7)'
    s["E7"].number_format = "0%"
    s.column_dimensions["A"].width = 16
    for col in "BCDEFG":
        s.column_dimensions[col].width = 18
    wb.save(path)


def recalc_check(path):
    """Fill a copy with sample answers, let LibreOffice recalculate, and compare to hand-counted expectations."""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "sample.xlsx")
    wb = load_workbook(path)
    log, q = wb["Log"], wb["Questions"]
    for cell, val in [("B2", "Acme"), ("B3", "invoicing tool"), ("B4", "a five-person agency"), ("B5", "teams that bill in several currencies"), ("B6", "Rival")]:
        q[cell] = val
    # engine, question, mentioned, wrong
    data = [
        ("ChatGPT", 1, "Yes", "No"), ("ChatGPT", 2, "No", "No"), ("ChatGPT", 3, "Yes", "Yes"), ("ChatGPT", 4, "Yes", "No"), ("ChatGPT", 5, "Yes", "No"),
        ("Claude", 1, "No", "No"), ("Claude", 2, "No", "No"), ("Claude", 3, "Yes", "No"), ("Claude", 4, "No", "No"), ("Claude", 5, "Yes", "Yes"),
        ("Perplexity", 1, "Yes", "No"), ("Perplexity", 2, "Yes", "No"), ("Perplexity", 3, "Yes", "No"), ("Perplexity", 4, "Yes", "No"), ("Perplexity", 5, "Yes", "No"),
    ]
    for i, (eng, n, ment, wrong) in enumerate(data, 2):
        log.cell(row=i, column=2, value=eng)
        log.cell(row=i, column=3, value=n)
        log.cell(row=i, column=6, value=ment)
        log.cell(row=i, column=9, value=wrong)
    wb.save(src)
    outdir = os.path.join(tmp, "out")
    os.makedirs(outdir)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", "xlsx", "--outdir", outdir, src],
                   check=True, capture_output=True, timeout=180)
    res = load_workbook(os.path.join(outdir, "sample.xlsx"), data_only=True)
    s, qs = res["Summary"], res["Questions"]
    expected = {  # row: (answers, unbranded, mentioned unbranded, rate, branded, wrong)
        4: (5, 2, 1, 0.5, 3, 1), 5: (5, 2, 0, 0.0, 3, 1), 6: (5, 2, 2, 1.0, 3, 0), 7: (15, 6, 3, 0.5, 9, 2)}
    problems = []
    for row, exp in expected.items():
        got = tuple(s.cell(row=row, column=c).value for c in range(2, 8))
        if got != exp:
            problems.append(f"Summary row {row}: expected {exp}, got {got}")
    q1 = qs["C9"].value
    if q1 != "What is the best invoicing tool for a five-person agency?":
        problems.append(f"Question 1 text unexpected: {q1!r}")
    qtexts = [qs.cell(row=r, column=3).value for r in range(9, 14)]
    shutil.rmtree(tmp, ignore_errors=True)
    return problems, qtexts


def robots_check():
    problems = []
    def parser(name):
        rp = urllib.robotparser.RobotFileParser()
        with open(os.path.join(HERE, name)) as f:
            rp.parse(f.read().splitlines())
        return rp
    default, optout = parser("ai-bots-robots.txt"), parser("ai-bots-robots-training-opt-out.txt")
    for agent in ["OAI-SearchBot", "Claude-SearchBot", "PerplexityBot"]:
        for label, rp in [("default", default), ("opt-out", optout)]:
            if not rp.can_fetch(agent, "https://example.com/pricing"):
                problems.append(f"{agent} should be allowed in the {label} file")
    for agent in ["GPTBot", "ClaudeBot", "Google-Extended"]:
        if default.can_fetch(agent, "https://example.com/"):
            pass  # default file leaves training open on purpose (rules are commented out)
        else:
            problems.append(f"{agent} should not be blocked in the default file (the opt-out lines are commented out)")
        if optout.can_fetch(agent, "https://example.com/"):
            problems.append(f"{agent} should be blocked in the opt-out file")
    return problems


def jsonld_check(vocab_path):
    vocab = json.load(open(vocab_path))["@graph"]
    classes, props = {}, {}
    def ids(x):
        if x is None: return []
        return [i["@id"] for i in (x if isinstance(x, list) else [x])]
    for node in vocab:
        t = node.get("@type")
        types = t if isinstance(t, list) else [t]
        if "rdfs:Class" in types:
            classes[node["@id"]] = ids(node.get("rdfs:subClassOf"))
        if "rdf:Property" in types:
            props[node["@id"]] = set(ids(node.get("schema:domainIncludes")))
    def ancestors(cls, seen=None):
        seen = seen or set()
        for p in classes.get(cls, []):
            if p not in seen:
                seen.add(p); ancestors(p, seen)
        return seen | {cls}
    problems = []
    def walk(obj, where):
        t = obj.get("@type")
        if t is None: return
        cid = f"schema:{t}"
        if cid not in classes:
            problems.append(f"{where}: unknown schema.org type {t}"); return
        valid = ancestors(cid)
        for key, val in obj.items():
            if key.startswith("@"): continue
            pid = f"schema:{key}"
            if pid not in props:
                problems.append(f"{where}: {key} is not a schema.org property")
            elif not (props[pid] & valid):
                problems.append(f"{where}: {key} is not valid on {t}")
            for v in (val if isinstance(val, list) else [val]):
                if isinstance(v, dict): walk(v, f"{where}.{key}")
    checked = []
    for name in ["organization", "product", "faqpage", "localbusiness"]:
        data = json.load(open(os.path.join(HERE, f"{name}.jsonld")))
        if data.get("@context") != "https://schema.org":
            problems.append(f"{name}: @context should be https://schema.org")
        walk(data, name)
        checked.append(name)
    return problems, checked


def main():
    build_xlsx(XLSX)
    results = []
    p, qtexts = recalc_check(XLSX)
    results.append(("Spreadsheet formulas recalculate to the hand-counted expectations (15 sample answers, 3 engines)", p))
    results.append(("robots.txt files: search bots allowed; training opt-out blocks GPTBot, ClaudeBot, Google-Extended", robots_check()))
    vocab = sys.argv[1] if len(sys.argv) > 1 else None
    if vocab and os.path.exists(vocab):
        p, checked = jsonld_check(vocab)
        results.append((f"JSON-LD ({', '.join(checked)}) parses, uses real schema.org types, and only valid properties", p))
        vocab_note = "validated against the current schema.org vocabulary file"
    else:
        vocab_note = "NOT validated against schema.org (vocabulary file not supplied)"
        results.append(("JSON-LD snippets", ["skipped: no vocabulary file supplied"]))
    lines = [f"# Verification of the free assets (run {date.today().isoformat()})", "",
             "Run `python3 plan/assets/make_and_verify.py <schemaorg-current-https.jsonld>` to reproduce.", ""]
    ok = True
    for title, problems in results:
        mark = "PASS" if not problems else "FAIL"
        ok &= not problems
        lines.append(f"- **{mark}** {title}")
        lines += [f"  - {x}" for x in problems]
    lines += ["", f"JSON-LD: {vocab_note}.", "",
              "Sample questions the template generated: " + " | ".join(qtexts), "",
              "## Not automated (needs a person)", "",
              "- Paste each JSON-LD file into Google's Rich Results Test and the schema.org validator. Google's tool cannot be driven from here.",
              "- Open the spreadsheet once in Excel or Google Sheets and confirm the drop-down lists show.",
              "- The robots.txt rules were tested with Python's parser. Crawlers can differ in edge cases, and the user-triggered fetchers may ignore robots.txt entirely."]
    open(os.path.join(HERE, "VERIFIED.md"), "w").write("\n".join(lines) + "\n")
    print("\n".join(lines))
    sys.exit(0 if ok else 1)


main()
